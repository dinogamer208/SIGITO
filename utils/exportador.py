"""
utils/exportador.py — Exportar reportes a PDF/Excel/CSV.

Responsable: Persona 5.

Qué debe hacer este archivo:
1. exportar_pdf(datos, ruta_salida) -> usando reportlab (o el skill de
   pdf si se genera desde un LLM, pero aquí es lógica pura de la app).
2. exportar_excel(datos, ruta_salida) -> usando openpyxl.
3. exportar_csv(datos, ruta_salida) -> csv de la librería estándar
   (no necesita dependencia extra).
4. `datos` debe ser una lista de dicts, tal como la devuelven las
   funciones de controllers/reportes_controller.py.

Esqueleto:
"""

import csv
import os


def exportar_csv(datos: list[dict], ruta_salida: str) -> None:
    if not datos:
        raise ValueError("No hay datos para exportar")
    with open(ruta_salida, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=datos[0].keys())
        writer.writeheader()
        writer.writerows(datos)


def importar_csv(ruta_entrada: str) -> list[dict]:
    with open(ruta_entrada, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def importar_excel(ruta_entrada: str) -> list[dict]:
    from openpyxl import load_workbook

    libro = load_workbook(ruta_entrada, read_only=True, data_only=True)
    hoja = libro.active

    filas = hoja.iter_rows(values_only=True)
    columnas = [str(c or "").strip() for c in next(filas)]

    datos = []
    for fila in filas:
        if all(valor is None for valor in fila):
            continue
        datos.append({
            columnas[i]: fila[i] for i in range(len(columnas))
        })
    return datos


def exportar_inventario_zip(datos: list[dict], ruta_salida: str) -> None:
    """
    Igual que exportar_csv, pero además mete en el mismo .zip una copia
    de cada foto referenciada en `foto_path` (carpeta "fotos/" dentro
    del zip), para poder llevar el inventario CON sus fotos a otra
    instalación de SIGITO (otra base de datos, otra PC).

    Dentro del zip, cada fila del CSV queda con foto_path apuntando a
    "fotos/<archivo>" (ruta relativa, portable) en vez de la ruta local
    original de esta PC.
    """
    import zipfile

    if not datos:
        raise ValueError("No hay datos para exportar")

    with zipfile.ZipFile(ruta_salida, "w", zipfile.ZIP_DEFLATED) as zf:
        fotos_ya_copiadas = set()
        filas_zip = []

        for fila in datos:
            fila_zip = dict(fila)
            ruta_foto = fila.get("foto_path")

            if ruta_foto and os.path.isfile(ruta_foto):
                nombre_archivo = os.path.basename(ruta_foto)
                if nombre_archivo not in fotos_ya_copiadas:
                    zf.write(ruta_foto, f"fotos/{nombre_archivo}")
                    fotos_ya_copiadas.add(nombre_archivo)
                fila_zip["foto_path"] = f"fotos/{nombre_archivo}"
            else:
                fila_zip["foto_path"] = ""

            filas_zip.append(fila_zip)

        columnas = list(filas_zip[0].keys())
        lineas = [",".join(columnas)]
        for fila in filas_zip:
            writer_fila = []
            for columna in columnas:
                valor = str(fila.get(columna, "")).replace('"', '""')
                if any(c in valor for c in (",", '"', "\n")):
                    valor = f'"{valor}"'
                writer_fila.append(valor)
            lineas.append(",".join(writer_fila))
        zf.writestr("inventario.csv", "\r\n".join(lineas))


def importar_inventario_zip(ruta_entrada: str, carpeta_fotos_destino: str) -> list[dict]:
    """
    Lee un .zip generado por exportar_inventario_zip: extrae las fotos a
    `carpeta_fotos_destino` (creándola si hace falta) y devuelve las
    filas del inventario con foto_path ya apuntando a esa carpeta local,
    listas para pasarle a inventario_controller.importar_articulos.
    """
    import zipfile

    with zipfile.ZipFile(ruta_entrada, "r") as zf:
        with zf.open("inventario.csv") as f:
            texto = f.read().decode("utf-8-sig").splitlines()
        filas = list(csv.DictReader(texto))

        os.makedirs(carpeta_fotos_destino, exist_ok=True)
        for nombre_en_zip in zf.namelist():
            if nombre_en_zip.startswith("fotos/") and not nombre_en_zip.endswith("/"):
                nombre_archivo = os.path.basename(nombre_en_zip)
                with zf.open(nombre_en_zip) as origen:
                    with open(os.path.join(carpeta_fotos_destino, nombre_archivo), "wb") as destino:
                        destino.write(origen.read())

    for fila in filas:
        if fila.get("foto_path"):
            nombre_archivo = os.path.basename(fila["foto_path"])
            fila["foto_path"] = os.path.join(carpeta_fotos_destino, nombre_archivo)

    return filas


def exportar_excel(datos: list[dict], ruta_salida: str) -> None:
    if not datos:
        raise ValueError("No hay datos para exportar")

    from openpyxl import Workbook
    from openpyxl.styles import Font

    columnas = list(datos[0].keys())

    libro = Workbook()
    hoja = libro.active

    hoja.append(columnas)
    for celda in hoja[1]:
        celda.font = Font(bold=True)

    for fila in datos:
        hoja.append([str(fila.get(col, "")) for col in columnas])

    for i, columna in enumerate(columnas, start=1):
        ancho = max(len(columna), *(len(str(fila.get(columna, ""))) for fila in datos)) + 2
        hoja.column_dimensions[hoja.cell(row=1, column=i).column_letter].width = ancho

    libro.save(ruta_salida)


def exportar_pdf(datos: list[dict], ruta_salida: str) -> None:
    if not datos:
        raise ValueError("No hay datos para exportar")

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

    columnas = list(datos[0].keys())
    filas = [[str(fila.get(col, "")) for col in columnas] for fila in datos]

    documento = SimpleDocTemplate(ruta_salida, pagesize=landscape(letter))
    tabla = Table([columnas] + filas, repeatRows=1)
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D4ED8")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))

    documento.build([tabla])


# ------------------------------------------------------------
# Reportes con varias secciones (ej. préstamos vencidos + ranking de
# más prestados). `secciones` es una lista de (titulo, datos), donde
# `datos` es una lista de dicts como en las funciones de arriba. Las
# secciones sin datos se omiten.
# ------------------------------------------------------------

def _secciones_con_datos(secciones):
    con_datos = [(titulo, datos) for titulo, datos in secciones if datos]
    if not con_datos:
        raise ValueError("No hay datos para exportar")
    return con_datos


def exportar_csv_secciones(secciones: list[tuple[str, list[dict]]], ruta_salida: str) -> None:
    """Un CSV no admite varias tablas: cada sección va una debajo de
    otra, con su título en una línea y una línea vacía entre ellas."""
    secciones = _secciones_con_datos(secciones)
    with open(ruta_salida, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        for i, (titulo, datos) in enumerate(secciones):
            if i:
                writer.writerow([])
            writer.writerow([titulo])
            columnas = list(datos[0].keys())
            writer.writerow(columnas)
            for fila in datos:
                writer.writerow([fila.get(col, "") for col in columnas])


def exportar_excel_secciones(secciones: list[tuple[str, list[dict]]], ruta_salida: str) -> None:
    """Una hoja por sección, con el título como nombre de la hoja."""
    secciones = _secciones_con_datos(secciones)

    from openpyxl import Workbook
    from openpyxl.styles import Font

    libro = Workbook()
    libro.remove(libro.active)

    for titulo, datos in secciones:
        # Excel limita el nombre de hoja a 31 caracteres y prohíbe algunos símbolos.
        nombre_hoja = "".join(c for c in titulo if c not in "[]:*?/\\")[:31]
        hoja = libro.create_sheet(nombre_hoja)

        columnas = list(datos[0].keys())
        hoja.append(columnas)
        for celda in hoja[1]:
            celda.font = Font(bold=True)

        for fila in datos:
            hoja.append([str(fila.get(col, "")) for col in columnas])

        for i, columna in enumerate(columnas, start=1):
            ancho = max(len(columna), *(len(str(fila.get(columna, ""))) for fila in datos)) + 2
            hoja.column_dimensions[hoja.cell(row=1, column=i).column_letter].width = ancho

    libro.save(ruta_salida)


def exportar_pdf_secciones(secciones: list[tuple[str, list[dict]]], ruta_salida: str) -> None:
    """Un solo PDF con el título de cada sección seguido de su tabla."""
    secciones = _secciones_con_datos(secciones)

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    estilos = getSampleStyleSheet()
    elementos = []

    for i, (titulo, datos) in enumerate(secciones):
        if i:
            elementos.append(Spacer(1, 18))
        elementos.append(Paragraph(titulo, estilos["Heading2"]))

        columnas = list(datos[0].keys())
        filas = [[str(fila.get(col, "")) for col in columnas] for fila in datos]

        tabla = Table([columnas] + filas, repeatRows=1)
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D4ED8")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        elementos.append(tabla)

    SimpleDocTemplate(ruta_salida, pagesize=landscape(letter)).build(elementos)
